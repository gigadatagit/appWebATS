from __future__ import annotations

import base64
import importlib.util
import io
import tempfile
import unicodedata
from copy import deepcopy
from pathlib import Path
from typing import Any


CHECK_ON = "\u2612"
CHECK_OFF = "\u2610"

_WORD_VALIDATION_DONE = False
_WORD_VALIDATION_ERROR = ""


def _b(v: bool) -> str:
    return CHECK_ON if v else CHECK_OFF


def _summarize_word_error(exc: Exception) -> str:
    raw = str(exc or "").strip()
    lowered = raw.lower()

    if "2147023584" in raw or "sesion de inicio" in lowered or "session" in lowered:
        return (
            "Word no puede iniciarse en la sesion actual de Windows. "
            "Ejecuta la app en una sesion interactiva con Word abierto/disponible."
        )
    if "word" in lowered and ("not installed" in lowered or "no se encontro" in lowered):
        return "Microsoft Word no esta disponible para docx2pdf."
    if "permission" in lowered or "permiso" in lowered or "access is denied" in lowered:
        return "Word/docx2pdf no tiene permisos para generar el PDF temporal."
    if raw:
        return raw
    return "Fallo desconocido al convertir con Word/docx2pdf."


def _normalize_text(value: Any) -> str:
    return str(value or "").strip()


def _normalize_key(value: str) -> str:
    text = _normalize_text(value).lower()
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


def _strip_data_uri(value: str | None) -> str:
    if not value:
        return ""
    if "," in value and value.lstrip().startswith("data:"):
        return value.split(",", 1)[1]
    return value.strip()


def _decode_base64_image(value: str | None) -> io.BytesIO | None:
    raw = _strip_data_uri(value)
    if not raw:
        return None
    try:
        return io.BytesIO(base64.b64decode(raw))
    except Exception:
        return None


def _iter_unique_cells(row):
    seen = set()
    out = []
    for cell in row.cells:
        key = id(cell._tc)
        if key not in seen:
            seen.add(key)
            out.append(cell)
    return out


def _clear_cell(cell) -> None:
    from docx.oxml import OxmlElement

    tc = cell._tc
    tc_pr = tc.tcPr
    for child in list(tc):
        if child is not tc_pr:
            tc.remove(child)
    p = OxmlElement("w:p")
    tc.append(p)


def _set_cell_vertical_center(cell) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    from docx.shared import Pt

    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _write_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size_pt: int = 9,
    center: bool = False,
    clear: bool = True,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    if clear:
        _clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run(text if text is not None else "")
    run.bold = bold
    run.font.size = Pt(size_pt)
    _set_cell_vertical_center(cell)


def _append_paragraph(cell, text: str, *, bold: bool = False, size_pt: int = 9, center: bool = False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    _set_cell_vertical_center(cell)
    return p


def _append_picture(cell, image_stream: io.BytesIO | None, *, width_cm: float = 3.0, centered: bool = True) -> bool:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    if image_stream is None:
        return False
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run()
    image_stream.seek(0)
    run.add_picture(image_stream, width=Cm(width_cm))
    _set_cell_vertical_center(cell)
    return True


def _insert_picture_before_signature_line(cell, image_stream: io.BytesIO | None, *, width_cm: float = 4.0) -> bool:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    if image_stream is None:
        return False

    signature_line = next((p for p in cell.paragraphs if "____" in p.text), None)
    if signature_line is not None:
        p = signature_line.insert_paragraph_before()
    else:
        _clear_cell(cell)
        p = cell.paragraphs[0]

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run()
    image_stream.seek(0)
    run.add_picture(image_stream, width=Cm(width_cm))

    if signature_line is None:
        _append_paragraph(cell, "____________________________________", size_pt=9, center=True)
        _append_paragraph(cell, "FIRMA", size_pt=9, center=True)

    _set_cell_vertical_center(cell)
    return True


def _set_cell_margins(cell, top=60, start=70, bottom=60, end=70):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def _normalize_table_padding(table):
    for row in table.rows:
        for cell in _iter_unique_cells(row):
            _set_cell_margins(cell)


def _table_text(table) -> str:
    parts: list[str] = []
    for row in table.rows:
        for cell in _iter_unique_cells(row):
            text = _normalize_key(cell.text)
            if text:
                parts.append(text)
    return " ".join(parts)


def _find_table_by_markers(doc, markers: list[str]):
    normalized_markers = [_normalize_key(marker) for marker in markers if _normalize_key(marker) != ""]
    if not normalized_markers:
        return None
    for table in doc.tables:
        full_text = _table_text(table)
        if all(marker in full_text for marker in normalized_markers):
            return table
    return None


def _get_template_table(doc, markers: list[str], fallback_index: int):
    table = _find_table_by_markers(doc, markers)
    if table is not None:
        return table
    if 0 <= fallback_index < len(doc.tables):
        return doc.tables[fallback_index]
    raise RuntimeError("No fue posible localizar la tabla requerida en la plantilla Word.")


def _remove_rows_from(table, start_index: int) -> None:
    while len(table.rows) > start_index:
        table._tbl.remove(table.rows[start_index]._tr)


def _safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(str(value).strip())
    except Exception:
        return default


def _extract_control_texts(raw_controls: Any) -> list[str]:
    result: list[str] = []
    if isinstance(raw_controls, list):
        for item in raw_controls:
            if isinstance(item, dict):
                text = _normalize_text(
                    item.get("control_aplicado")
                    or item.get("control_nombre")
                    or item.get("control")
                    or item.get("nombre")
                )
            else:
                text = _normalize_text(item)
            if text and text not in result:
                result.append(text)
        return result

    single = _normalize_text(raw_controls)
    return [single] if single else []


def _extract_step_peligros(step_item: dict[str, Any]) -> list[dict[str, Any]]:
    raw_peligros = step_item.get("peligros")
    if isinstance(raw_peligros, list) and raw_peligros:
        out: list[dict[str, Any]] = []
        for item in raw_peligros:
            if not isinstance(item, dict):
                continue
            peligro_numero = _safe_int(item.get("peligro_numero"), 0)
            peligro_nombre = _normalize_text(item.get("peligro_nombre") or item.get("nombre"))
            controles = _extract_control_texts(item.get("controles", []))
            if peligro_numero <= 0 and peligro_nombre == "" and not controles:
                continue
            out.append(
                {
                    "peligro_numero": peligro_numero,
                    "peligro_nombre": peligro_nombre,
                    "controles": controles,
                }
            )
        return out

    # fallback plano legado
    danger_parts = [
        part.strip()
        for part in _normalize_text(step_item.get("peligros_asociados")).split("|")
        if part.strip() != ""
    ]
    control_parts = [
        part.strip()
        for part in _normalize_text(step_item.get("controles_aplicados")).split("|")
        if part.strip() != ""
    ]
    if not danger_parts and not control_parts:
        return []
    return [
        {
            "peligro_numero": 0,
            "peligro_nombre": " | ".join(danger_parts),
            "controles": control_parts,
        }
    ]


def _write_controls_as_lines(cell, controls: list[str], *, size_pt: int = 9):
    bullet = "\u2022"
    if not controls:
        _write_cell_text(cell, "", size_pt=size_pt)
        return
    _write_cell_text(cell, f"{bullet} {controls[0]}", size_pt=size_pt)
    for control in controls[1:]:
        _append_paragraph(cell, f"{bullet} {control}", size_pt=size_pt)


def _map_context_to_template_data(context: dict[str, Any]) -> dict[str, Any]:
    ats = context.get("ats", {})
    apoyos = context.get("apoyos", [])
    certificados = context.get("certificados", [])
    pasos = context.get("pasos", [])
    trabajadores = context.get("trabajadores", [])
    firmas_finales = context.get("firmas_finales", [])
    peligros = context.get("peligros", [])

    apoyo_flags = {
        "programa": False,
        "procedimiento": False,
        "instructivo": False,
        "ordenServicio": False,
        "planEmergencia": False,
        "procMantenimiento": False,
        "otro": False,
        "otroCual": "",
    }
    for item in apoyos:
        code = _normalize_key(str(item.get("codigo") or ""))
        name = _normalize_key(str(item.get("nombre") or ""))
        extra = _normalize_text(item.get("descripcion_otro"))
        if code == "programa" or "programa" in name:
            apoyo_flags["programa"] = True
        elif code in {"proc_mantenimiento", "procedimiento_mantenimiento"} or "procedimiento de mantenimiento" in name:
            apoyo_flags["procMantenimiento"] = True
        elif code == "procedimiento" or "procedimiento" in name:
            apoyo_flags["procedimiento"] = True
        elif code == "instructivo" or "instructivo" in name:
            apoyo_flags["instructivo"] = True
        elif code == "orden_servicio" or "orden de servicio" in name:
            apoyo_flags["ordenServicio"] = True
        elif code == "plan_emergencia" or "plan de emergencia" in name:
            apoyo_flags["planEmergencia"] = True
        elif code in {"otro", "otros"} or "otro" in name:
            apoyo_flags["otro"] = True
            if extra:
                apoyo_flags["otroCual"] = extra

    cert_flags = {
        "alturas": False,
        "electricos": False,
        "espaciosConfinados": False,
        "trabajosCaliente": False,
        "otro": False,
        "otroCual": "",
    }
    for item in certificados:
        code = _normalize_key(str(item.get("codigo") or ""))
        name = _normalize_key(str(item.get("nombre") or ""))
        extra = _normalize_text(item.get("descripcion_otro"))
        if code == "alturas" or "altura" in name:
            cert_flags["alturas"] = True
        elif code == "electricos" or "electr" in name:
            cert_flags["electricos"] = True
        elif code == "espacios_confinados" or "confinad" in name:
            cert_flags["espaciosConfinados"] = True
        elif code in {"trabajos_en_caliente", "trabajos_caliente"} or "caliente" in name:
            cert_flags["trabajosCaliente"] = True
        elif code in {"otro", "otros"} or "otro" in name:
            cert_flags["otro"] = True
            if extra:
                cert_flags["otroCual"] = extra

    otros_peligros = ""
    for p in peligros:
        nombre = _normalize_key(str(p.get("nombre") or ""))
        if "otro" in nombre and _normalize_text(p.get("descripcion_otro")):
            otros_peligros = _normalize_text(p.get("descripcion_otro"))
            break

    pasos_template: list[dict[str, Any]] = []
    for item in pasos:
        peligros_detalle = _extract_step_peligros(item)
        peligros_asociados = _normalize_text(item.get("peligros_asociados"))
        if peligros_asociados == "" and peligros_detalle:
            peligros_asociados = " | ".join(
                [
                    f"P{int(p.get('peligro_numero') or 0)} - {str(p.get('peligro_nombre') or '')}"
                    if int(p.get("peligro_numero") or 0) > 0
                    else str(p.get("peligro_nombre") or "")
                    for p in peligros_detalle
                    if str(p.get("peligro_nombre") or "").strip() != "" or int(p.get("peligro_numero") or 0) > 0
                ]
            )
        controles_aplicados = _normalize_text(item.get("controles_aplicados"))
        if controles_aplicados == "" and peligros_detalle:
            controles_aplicados = " | ".join(
                list(
                    dict.fromkeys(
                        [
                            str(control or "").strip()
                            for peligro in peligros_detalle
                            for control in peligro.get("controles", [])
                            if str(control or "").strip() != ""
                        ]
                    )
                )
            )

        pasos_template.append(
            {
                "numeroPaso": int(item.get("numero_paso") or 0),
                "descripcionPasosASeguir": _normalize_text(item.get("descripcion_paso")),
                "peligros": peligros_detalle,
                "peligrosAsociados": peligros_asociados,
                "controlesARealizar": controles_aplicados,
            }
        )

    trabajadores_template = [
        {
            "nombreTrabajador": _normalize_text(item.get("nombre_trabajador")),
            "numeroDocumentoTrabajador": _normalize_text(item.get("numero_documento")),
            "cargoTrabajador": _normalize_text(item.get("cargo_trabajador")),
            "firmaTrabajadorBase64": _normalize_text(item.get("firma_base64")),
        }
        for item in trabajadores
    ]

    firmas_template = [
        {
            "tipoFirma": _normalize_text(item.get("firma_tipo_codigo")).upper(),
            "nombreCompleto": _normalize_text(item.get("nombre_completo")),
            "cargo": _normalize_text(item.get("cargo")),
            "firmaBase64": _normalize_text(item.get("firma_base64")),
        }
        for item in firmas_finales
    ]

    tipo_ats = _normalize_key(str(ats.get("tipo_ats_codigo") or ats.get("tipo_ats") or "")).upper()

    return {
        "ats": {
            "empresaOPersonaEjecuta": _normalize_text(ats.get("empresa_persona_ejecuta")),
            "fechaElaboracion": _normalize_text(ats.get("fecha_elaboracion")),
            "ciudad": _normalize_text(ats.get("ciudad")),
            "areaOLugar": _normalize_text(ats.get("area_lugar")),
            "numeroAts": _normalize_text(ats.get("numero_ats")),
            "tipoAts": tipo_ats,
            "duracionActividad": _normalize_text(ats.get("duracion_actividad")),
            "actividadAltoRiesgo": _normalize_text(ats.get("actividad_alto_riesgo")).upper() == "SI",
            "descripcionActividad": _normalize_text(ats.get("descripcion_actividad")),
            "informacionApoyo": apoyo_flags,
            "certificadosAsociados": cert_flags,
            "otrosPeligros": otros_peligros,
        },
        "pasos": pasos_template,
        "trabajadores": trabajadores_template,
        "observaciones": _normalize_text(context.get("observaciones")),
        "firmasFinales": firmas_template,
    }


def _fill_identificacion_general(doc, data: dict[str, Any]) -> None:
    table = _get_template_table(
        doc,
        markers=["IDENTIFICACIÓN GENERAL", "Fecha Elaboración"],
        fallback_index=0,
    )
    ats = data["ats"]
    info = ats.get("informacionApoyo", {})
    cert = ats.get("certificadosAsociados", {})

    _write_cell_text(table.cell(1, 8), ats.get("empresaOPersonaEjecuta", ""), size_pt=10)
    _write_cell_text(table.cell(1, 15), ats.get("fechaElaboracion", ""), size_pt=10, center=True)
    _write_cell_text(table.cell(2, 1), ats.get("ciudad", ""), size_pt=10, center=True)
    _write_cell_text(table.cell(2, 10), ats.get("areaOLugar", ""), size_pt=10, center=True)
    _write_cell_text(table.cell(3, 2), ats.get("numeroAts", ""), size_pt=10, center=True)
    _write_cell_text(
        table.cell(3, 3),
        f"Nuevo: {_b(str(ats.get('tipoAts', '')).upper() == 'NUEVO')}      Revisi\u00f3n: {_b(str(ats.get('tipoAts', '')).upper() == 'REVISION')}",
        size_pt=9,
        center=True,
    )
    _write_cell_text(table.cell(3, 10), ats.get("duracionActividad", ""), size_pt=10, center=True)
    _write_cell_text(
        table.cell(3, 11),
        f"Actividad de Alto Riesgo:   SI {_b(bool(ats.get('actividadAltoRiesgo', False)))}      NO {_b(not bool(ats.get('actividadAltoRiesgo', False)))}",
        size_pt=9,
        center=True,
    )
    _write_cell_text(table.cell(4, 6), ats.get("descripcionActividad", ""), size_pt=10)

    apoyo_lines = [
        f"Programa {_b(bool(info.get('programa', False)))}     Procedimiento {_b(bool(info.get('procedimiento', False)))}     "
        f"Instructivo {_b(bool(info.get('instructivo', False)))}     Orden de Servicio {_b(bool(info.get('ordenServicio', False)))}",
        "",
        f"Plan de Emergencia {_b(bool(info.get('planEmergencia', False)))}     Procedimiento de mantenimiento {_b(bool(info.get('procMantenimiento', False)))}",
    ]
    _write_cell_text(table.cell(5, 4), "\n".join(apoyo_lines), size_pt=9)
    _write_cell_text(
        table.cell(5, 14),
        f"Otros {_b(bool(info.get('otro', False)))}\n\u00bfCu\u00e1l?: {info.get('otroCual', '')}",
        size_pt=9,
    )

    cert_text = (
        f"Alturas {_b(bool(cert.get('alturas', False)))}     "
        f"Electricos {_b(bool(cert.get('electricos', False)))}     "
        f"Esp. Confinados {_b(bool(cert.get('espaciosConfinados', False)))}     "
        f"Trabajos en Caliente {_b(bool(cert.get('trabajosCaliente', False)))}"
    )
    _write_cell_text(table.cell(6, 5), cert_text, size_pt=9)
    _write_cell_text(
        table.cell(6, 13),
        f"Otros {_b(bool(cert.get('otro', False)))}\n\u00bfCu\u00e1l?: {cert.get('otroCual', '')}",
        size_pt=9,
    )

    _normalize_table_padding(table)


def _fill_otros_peligros(doc, data: dict[str, Any]) -> None:
    text = data.get("ats", {}).get("otrosPeligros", "")
    if not text:
        return
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Otros:":
            if paragraph.runs:
                paragraph.runs[0].text = f"Otros: {text}"
            else:
                paragraph.add_run(f"Otros: {text}")
            break


def _fill_steps_table(doc, data: dict[str, Any]) -> None:
    table = _get_template_table(
        doc,
        markers=["Descripción de los Pasos Para Seguir", "PELIGROS ASOCIADOS", "Controles a Realizar"],
        fallback_index=2,
    )
    pasos = data.get("pasos", [])
    _normalize_table_padding(table)

    if len(table.rows) < 2:
        raise RuntimeError("La tabla de pasos en la plantilla no contiene fila prototipo para subfilas.")

    prototype_tr = deepcopy(table.rows[1]._tr)
    _remove_rows_from(table, 1)

    if not pasos:
        table._tbl.append(deepcopy(prototype_tr))
        row = table.rows[-1]
        _write_cell_text(row.cells[0], "1", size_pt=10, center=True)
        _write_cell_text(row.cells[1], "Sin pasos registrados.", size_pt=9)
        _write_cell_text(row.cells[2], "", size_pt=9)
        _write_cell_text(row.cells[3], "", size_pt=9)
        return

    for idx, paso in enumerate(pasos, start=1):
        table._tbl.append(deepcopy(prototype_tr))
        header_row = table.rows[-1]

        numero_paso = int(paso.get("numeroPaso") or idx)
        descripcion = _normalize_text(paso.get("descripcionPasosASeguir"))
        _write_cell_text(header_row.cells[0], str(numero_paso), size_pt=10, center=True, bold=True)
        _write_cell_text(header_row.cells[1], descripcion, size_pt=9, bold=True)
        _write_cell_text(header_row.cells[2], "", size_pt=9)
        _write_cell_text(header_row.cells[3], "", size_pt=9)
        try:
            merged = header_row.cells[1].merge(header_row.cells[3])
            _write_cell_text(merged, descripcion, size_pt=9, bold=True)
        except Exception:
            pass

        peligros: list[dict[str, Any]] = [dict(item) for item in paso.get("peligros", [])]
        if not peligros:
            table._tbl.append(deepcopy(prototype_tr))
            empty_row = table.rows[-1]
            _write_cell_text(empty_row.cells[0], "", size_pt=9, center=True)
            _write_cell_text(empty_row.cells[1], "", size_pt=9)
            _write_cell_text(empty_row.cells[2], "Sin peligros asociados.", size_pt=9)
            _write_cell_text(empty_row.cells[3], "", size_pt=9)
            continue

        for peligro in peligros:
            table._tbl.append(deepcopy(prototype_tr))
            subrow = table.rows[-1]
            _write_cell_text(subrow.cells[0], "", size_pt=9, center=True)
            _write_cell_text(subrow.cells[1], "", size_pt=9)

            peligro_numero = int(peligro.get("peligro_numero") or 0)
            peligro_nombre = _normalize_text(peligro.get("peligro_nombre"))
            peligro_texto = (
                f"P{peligro_numero} - {peligro_nombre}"
                if peligro_numero > 0 and peligro_nombre != ""
                else peligro_nombre
            )
            _write_cell_text(subrow.cells[2], peligro_texto, size_pt=9)

            controles = [str(control or "").strip() for control in peligro.get("controles", []) if str(control or "").strip() != ""]
            _write_controls_as_lines(subrow.cells[3], controles, size_pt=9)


def _fill_trabajadores_table(doc, data: dict[str, Any]) -> None:
    table = _get_template_table(
        doc,
        markers=["TRABAJADORES RELACIONADOS CON LA ACTIVIDAD", "DOCUMENTO DE ID", "FIRMA"],
        fallback_index=3,
    )
    trabajadores = data.get("trabajadores", [])
    _normalize_table_padding(table)

    while len(table.rows) - 2 < len(trabajadores):
        table._tbl.append(deepcopy(table.rows[-1]._tr))

    for idx, trabajador in enumerate(trabajadores, start=2):
        row = table.rows[idx]
        _write_cell_text(row.cells[0], trabajador.get("nombreTrabajador", ""), size_pt=9)
        _write_cell_text(row.cells[1], trabajador.get("numeroDocumentoTrabajador", ""), size_pt=9, center=True)
        _write_cell_text(row.cells[2], trabajador.get("cargoTrabajador", ""), size_pt=9)
        _clear_cell(row.cells[3])
        img = _decode_base64_image(trabajador.get("firmaTrabajadorBase64"))
        if not _append_picture(row.cells[3], img, width_cm=2.8):
            _write_cell_text(row.cells[3], "", size_pt=9, center=True, clear=True)

    for row_idx in range(len(trabajadores) + 2, len(table.rows)):
        row = table.rows[row_idx]
        for col in range(4):
            _write_cell_text(row.cells[col], "", size_pt=9, center=(col in (1, 3)))


def _fill_observaciones(doc, data: dict[str, Any]) -> None:
    table = _get_template_table(
        doc,
        markers=["Observaciones"],
        fallback_index=4,
    )
    _write_cell_text(table.cell(1, 0), data.get("observaciones", ""), size_pt=10)


def _fill_firmas_finales(doc, data: dict[str, Any]) -> None:
    table = _get_template_table(
        doc,
        markers=["AUTORIZA ACTIVIDAD SST", "SUPERVISA ACTIVIDAD SST", "EJECUTA ACTIVIDAD"],
        fallback_index=5,
    )
    _normalize_table_padding(table)
    roles = {"AUTORIZA": 0, "SUPERVISA": 1, "EJECUTA": 2}
    firmas = {item.get("tipoFirma", "").upper(): item for item in data.get("firmasFinales", [])}

    for role, col in roles.items():
        item = firmas.get(role, {})
        cell0 = table.cell(0, col)
        _insert_picture_before_signature_line(cell0, _decode_base64_image(item.get("firmaBase64")), width_cm=4.0)
        cell1 = table.cell(1, col)
        _write_cell_text(cell1, f"NOMBRE: {item.get('nombreCompleto', '')}", size_pt=9)
        _append_paragraph(cell1, f"CARGO: {item.get('cargo', '')}", size_pt=9)
        label = {
            "AUTORIZA": "AUTORIZA ACTIVIDAD SST",
            "SUPERVISA": "SUPERVISA ACTIVIDAD SST",
            "EJECUTA": "EJECUTA ACTIVIDAD",
        }[role]
        _write_cell_text(table.cell(2, col), label, size_pt=9, center=True)


def generate_docx_from_template(context: dict[str, Any], template_path: Path, output_docx_path: Path) -> Path:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("Falta dependencia python-docx para usar plantilla DOCX.") from exc

    payload = _map_context_to_template_data(context)
    doc = Document(str(template_path))
    _fill_identificacion_general(doc, payload)
    _fill_otros_peligros(doc, payload)
    _fill_steps_table(doc, payload)
    _fill_trabajadores_table(doc, payload)
    _fill_observaciones(doc, payload)
    _fill_firmas_finales(doc, payload)
    output_docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_docx_path))
    return output_docx_path


def _validate_word_dependencies():
    missing: list[str] = []
    if importlib.util.find_spec("docx") is None:
        missing.append("python-docx")
    if importlib.util.find_spec("docx2pdf") is None:
        missing.append("docx2pdf")
    if missing:
        raise RuntimeError("Faltan dependencias requeridas: " + ", ".join(missing))


def _run_word_smoke_test():
    from docx import Document
    from docx2pdf import convert  # type: ignore

    with tempfile.TemporaryDirectory(prefix="ats_word_validate_") as tmp_dir:
        tmp_path = Path(tmp_dir)
        docx_path = tmp_path / "word_validation.docx"
        pdf_path = tmp_path / "word_validation.pdf"

        doc = Document()
        doc.add_paragraph("ATS Word validation")
        doc.save(str(docx_path))

        convert(str(docx_path), str(pdf_path))
        if not pdf_path.exists() or pdf_path.stat().st_size <= 0:
            raise RuntimeError("Word/docx2pdf no genero PDF en la validacion previa.")


def ensure_word_conversion_ready():
    global _WORD_VALIDATION_DONE, _WORD_VALIDATION_ERROR

    if _WORD_VALIDATION_DONE:
        if _WORD_VALIDATION_ERROR:
            raise RuntimeError(_WORD_VALIDATION_ERROR)
        return

    try:
        _validate_word_dependencies()
        _run_word_smoke_test()
        _WORD_VALIDATION_ERROR = ""
    except Exception as exc:
        _WORD_VALIDATION_ERROR = _summarize_word_error(exc)
        raise RuntimeError(_WORD_VALIDATION_ERROR) from exc
    finally:
        _WORD_VALIDATION_DONE = True


def _convert_docx_to_pdf(docx_path: Path, output_pdf_path: Path):
    try:
        from docx2pdf import convert  # type: ignore

        output_pdf_path.parent.mkdir(parents=True, exist_ok=True)
        convert(str(docx_path), str(output_pdf_path))
        if output_pdf_path.exists():
            return
    except Exception as exc:
        raise RuntimeError(_summarize_word_error(exc)) from exc

    raise RuntimeError("Word/docx2pdf no genero el archivo PDF esperado.")


def generate_pdf_from_template(context: dict[str, Any], template_path: Path, output_pdf_path: Path) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"No existe la plantilla DOCX: {template_path}")

    ensure_word_conversion_ready()

    output_pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ats_word_build_") as tmp_dir:
        tmp_docx = Path(tmp_dir) / "ats_render.docx"
        generate_docx_from_template(context, template_path, tmp_docx)
        _convert_docx_to_pdf(tmp_docx, output_pdf_path)

    if not output_pdf_path.exists():
        raise RuntimeError("No se genero archivo PDF desde plantilla Word.")
    return output_pdf_path
